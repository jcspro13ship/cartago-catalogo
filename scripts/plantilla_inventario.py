"""Genera docs/Cartago_Inventario_Plantilla.xlsx: la plantilla del Google Sheet
del inventario (Léeme, Productos, Variantes, Imagenes, Marcas, Configuracion,
Listas). Se sube a Drive y Google la convierte en Sheet.

Uso: python scripts/plantilla_inventario.py docs/Cartago_Inventario_Plantilla.xlsx
Requiere openpyxl.
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VINO = '771541'
EJEMPLO = 'F4E9ED'
ENTRADA = 'FFF7D6'
F = 'Arial'
FILAS = 150  # el catálogo es de ~100 productos

CATEGORIAS = {
    'Monturas': ['Sillas', 'Tereques', 'Galápagos', 'Kids'],
    'Equino': ['Cabeza', 'Alfombras', 'Alforjas', 'Correaje', 'Tarabas'],
    'Jinete': ['Sombreros', 'Marroquinería', 'Bisutería'],
    'Cuidado del caballo': ['Aseo', 'Suplementos', 'Accesorios'],
    'Caninos': ['Pecheras', 'Collares', 'Traíllas', 'Bozales'],
}
MARCAS = ['Cartago', 'San Fermín', 'Mesacé', 'Al Torno', 'Sybec', 'Natural V.', 'Cheval']

encabezado = Font(name=F, bold=True, color='FFFFFF', size=10)
relleno_enc = PatternFill('solid', fgColor=VINO)
cuerpo = Font(name=F, size=10)
relleno_ej = PatternFill('solid', fgColor=EJEMPLO)
calculada = Font(name=F, size=10, color='6E5F66', italic=True)


def tabla(ws, columnas, notas=None, congelar='B2'):
    notas = notas or {}
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=i, value=nombre)
        c.font = encabezado
        c.fill = relleno_enc
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = ancho
        if nombre in notas:
            c.comment = Comment(notas[nombre], 'JCS Emprende')
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = congelar


def fila_ejemplo(ws, valores, fila=2):
    for i, v in enumerate(valores, start=1):
        if v is None:
            continue
        c = ws.cell(row=fila, column=i, value=v)
        c.font = cuerpo
        c.fill = relleno_ej
        c.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str) and len(v) > 30)


def lista(ws, col, fuente, filas=500):
    dv = DataValidation(type='list', formula1=fuente, allow_blank=True)
    dv.error = 'Elige un valor de la lista.'
    dv.errorTitle = 'Valor no válido'
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{filas}')


wb = Workbook()

# ---------- Léeme ----------
ws = wb.active
ws.title = 'Léeme'
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 112
ws.sheet_view.showGridLines = False
lineas = [
    ('t', 'Inventario del catálogo web · Talabartería Cartago'),
    ('s', 'Esta hoja alimenta la página web. Lo que cambien aquí aparece en el sitio en máximo 5 minutos.'),
    ('', ''),
    ('h', 'Pestañas'),
    ('p', 'Productos: una fila por producto (nombre, categoría, subcategoría, marca, descripción, precio).'),
    ('p', 'Variantes: una fila por cada color de cada producto, con su stock.'),
    ('p', 'Imagenes: una fila por foto de cada producto, en orden. La foto 1 es la portada.'),
    ('p', 'Marcas: las marcas que vende Cartago y el link de su logo (opcional).'),
    ('p', 'Configuracion: la tasa del dólar. Se actualiza cada 8 días.'),
    ('p', 'Listas: categorías y subcategorías permitidas.'),
    ('', ''),
    ('h', 'Cómo llenar'),
    ('p', '1. La fila 2 de cada pestaña (en rosado) es un ejemplo. Bórrenla cuando carguen el primer producto real: si no, aparece en la página.'),
    ('p', '2. Cada producto tiene un id corto y único: una letra y un número (m1, m2 para monturas; e1 para equino; j1 para jinete). Nunca reutilicen un id.'),
    ('p', '3. Categoría y subcategoría se eligen de la lista. Ej.: Equino > Cabeza para aperos, jáquimas, riendas y frenos.'),
    ('p', '4. Precio en pesos sin puntos ni signo: 85000. Precio 0 o vacío = "Precio a consultar" (silla mexicana, frenos a la medida, Chaparral personalizada...).'),
    ('p', '5. precio_usd es opcional: solo si quieren un precio fijo en dólares. Si queda vacío, se calcula con la tasa de Configuracion.'),
    ('p', '6. activo = NO esconde el producto sin borrarlo. destacado = SI lo muestra en la página de inicio.'),
    ('p', '7. personalizable = SI muestra un campo para que el cliente escriba el bordado (nombre del caballo, criadero o iniciales). nota_personalizacion es la pista que ve el cliente.'),
    ('p', '8. garantia es opcional: un texto que aparece resaltado en la ficha. Ej.: la garantía de cambio de los frenos.'),
    ('p', '9. En Variantes, si el producto no maneja colores, dejen una sola fila con el color vacío. Con stock 0 el color aparece como agotado.'),
    ('p', f'10. Las columnas grises en cursiva (stock_total, fotos) se calculan solas. Vienen listas para {FILAS} productos; si pasan de ahí, copien la última fórmula hacia abajo.'),
    ('', ''),
    ('h', 'Fotos'),
    ('p', 'Suban las fotos a la carpeta "Fotos de productos" de este mismo Drive, nombradas con el id del producto y un número: e2-1.jpg, e2-2.jpg.'),
    ('p', 'JCS Emprende las optimiza y escribe el link en la pestaña Imagenes. Buena luz, fondo limpio y al menos 1200 px de lado.'),
    ('', ''),
    ('h', 'Al vender'),
    ('p', 'Cuando el cliente pague, resten las unidades vendidas en la columna stock de la pestaña Variantes.'),
    ('', ''),
    ('n', 'No cambien los nombres de las pestañas ni de los encabezados: la página los usa para leer los datos.'),
]
for r, (tipo, texto) in enumerate(lineas, start=2):
    c = ws.cell(row=r, column=2, value=texto)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if tipo == 't':
        c.font = Font(name=F, bold=True, size=16, color=VINO)
        ws.row_dimensions[r].height = 28
    elif tipo == 's':
        c.font = Font(name=F, size=11, color='6E5F66')
    elif tipo == 'h':
        c.font = Font(name=F, bold=True, size=12, color=VINO)
        ws.row_dimensions[r].height = 22
    elif tipo == 'n':
        c.font = Font(name=F, bold=True, size=10, color='B42318')
    else:
        c.font = cuerpo

# ---------- Productos ----------
ws = wb.create_sheet('Productos')
columnas = [
    ('id', 8), ('nombre', 28), ('categoria', 20), ('subcategoria', 16), ('marca', 14), ('descripcion', 56),
    ('precio', 12), ('precio_usd', 11), ('activo', 8), ('destacado', 10), ('personalizable', 13),
    ('nota_personalizacion', 40), ('garantia', 40), ('stock_total', 11), ('fotos', 7),
]
tabla(ws, columnas, {
    'id': 'Código corto y único: letra + número (m1, e1, j1). No se repite ni se cambia.',
    'categoria': 'Elige de la lista (pestaña Listas).',
    'subcategoria': 'Elige de la lista. Debe corresponder a la categoría (ver pestaña Listas).',
    'marca': 'Elige de la lista (pestaña Marcas). Cada producto muestra el logo de su marca.',
    'precio': 'Pesos colombianos, sin puntos ni signo. 0 o vacío = Precio a consultar.',
    'precio_usd': 'Opcional. Precio fijo en dólares. Vacío = se convierte con la tasa de Configuracion.',
    'activo': 'SI = se ve en la página. NO = escondido.',
    'destacado': 'SI = aparece en Productos destacados de la página de inicio.',
    'personalizable': 'SI = el cliente puede escribir el bordado o el detalle que quiere.',
    'nota_personalizacion': 'Texto de ayuda que ve el cliente en el campo de personalización.',
    'garantia': 'Opcional. Texto resaltado en la ficha del producto.',
    'stock_total': 'Se calcula solo: suma del stock de todos los colores en Variantes.',
    'fotos': 'Se calcula solo: cuántas fotos tiene el producto en Imagenes.',
})
fila_ejemplo(ws, [
    'e2', 'Freno a la medida', 'Equino', 'Cabeza', 'Cartago', 'Freno fabricado bajo medida para tu caballo.',
    0, None, 'SI', 'SI', 'NO', None,
    'Garantía de cambio: lo ensayas y, si no se le acomoda a tu caballo, te lo cambiamos sin costo adicional.',
])
for r in range(2, FILAS + 2):
    for col, formula in (
        ('N', f'=IF(A{r}="","",SUMIF(Variantes!A:A,A{r},Variantes!D:D))'),
        ('O', f'=IF(A{r}="","",COUNTIF(Imagenes!A:A,A{r}))'),
    ):
        c = ws[f'{col}{r}']
        c.value = formula
        c.font = calculada
ws['G2'].number_format = '#,##0'
lista(ws, 'C', '=Listas!$A$2:$A$20')
lista(ws, 'D', '=Listas!$C$2:$C$60')
lista(ws, 'E', '=Marcas!$A$2:$A$50')
for col in 'IJK':
    lista(ws, col, '"SI,NO"')

# ---------- Variantes ----------
ws = wb.create_sheet('Variantes')
tabla(ws, [('producto_id', 14), ('sku', 16), ('color', 20), ('stock', 10)], {
    'producto_id': 'El id del producto, tal como está en Productos.',
    'sku': 'Código único de esta variante: id + 3 letras del color (j2-NEG). Sin colores: el mismo id.',
    'color': 'Nombre del color. Vacío si el producto no maneja colores.',
    'stock': 'Unidades disponibles de este color. 0 = agotado.',
})
fila_ejemplo(ws, ['e2', 'e2', None, 10])
lista(ws, 'A', '=Productos!$A$2:$A$500')

# ---------- Imagenes ----------
ws = wb.create_sheet('Imagenes')
tabla(ws, [('producto_id', 14), ('orden', 8), ('url', 70), ('archivo', 22)], {
    'orden': '1 = foto de portada; luego 2, 3...',
    'url': 'Link de la foto en Cloudinary (lo escribe JCS Emprende).',
    'archivo': 'Nombre del archivo en la carpeta de Drive, para ubicarlo (e2-1.jpg).',
})
fila_ejemplo(ws, ['e2', 1, 'https://res.cloudinary.com/.../e2-1.jpg', 'e2-1.jpg'])
lista(ws, 'A', '=Productos!$A$2:$A$500')

# ---------- Marcas ----------
ws = wb.create_sheet('Marcas')
tabla(ws, [('marca', 18), ('logo_url', 70)], {
    'marca': 'Nombre de la marca tal como se escribe en Productos.',
    'logo_url': 'Opcional. Link del logo. Cartago, San Fermín y Mesacé ya vienen incluidos en la página.',
}, congelar='A2')
for i, m in enumerate(MARCAS, start=2):
    ws.cell(row=i, column=1, value=m).font = cuerpo

# ---------- Configuracion ----------
ws = wb.create_sheet('Configuracion')
tabla(ws, [('clave', 14), ('valor', 12), ('descripcion', 70)], congelar='A2')
for i, v in enumerate(['tasa_usd', 3000, 'Pesos por dólar para mostrar precios en USD. Se actualiza cada 8 días.'], start=1):
    ws.cell(row=2, column=i, value=v).font = cuerpo
ws['B2'].fill = PatternFill('solid', fgColor=ENTRADA)
ws['B2'].number_format = '#,##0'

# ---------- Listas ----------
ws = wb.create_sheet('Listas')
tabla(ws, [('categorias', 24), ('', 3), ('subcategorias', 20), ('de la categoría', 24)], {
    'categorias': 'Categorías que se pueden elegir en Productos.',
    'subcategorias': 'Subcategorías que se pueden elegir en Productos.',
}, congelar='A2')
for i, cat in enumerate(CATEGORIAS, start=2):
    ws.cell(row=i, column=1, value=cat).font = cuerpo
fila = 2
for cat, subs in CATEGORIAS.items():
    for sub in subs:
        ws.cell(row=fila, column=3, value=sub).font = cuerpo
        ws.cell(row=fila, column=4, value=cat).font = calculada
        fila += 1

wb.save(sys.argv[1])
print('Plantilla guardada en', sys.argv[1])
